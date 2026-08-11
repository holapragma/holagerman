"""
Genera la propuesta comercial de German CRM en formato .docx.
Estilo: minimalista, mucho espacio en blanco, inspirado en Apple / Stripe / Linear / Notion.
Nunca editar el .docx a mano — todo cambio se hace en este script y se re-genera.
"""

import os
from docx import Document
from docx.shared import Mm, Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

ASSETS = os.path.join(os.path.dirname(__file__), "assets")
OUT = os.path.join(os.path.dirname(__file__), "propuesta_crm_comercial_v1.docx")

FONT = "Helvetica Neue"
GREEN = RGBColor(0x0F, 0x51, 0x32)
DARK = RGBColor(0x10, 0x18, 0x28)
MUTED = RGBColor(0x6B, 0x72, 0x80)
LIGHT_MUTED = RGBColor(0x9C, 0xA3, 0xAF)
WARN_BG_TEXT = RGBColor(0x8A, 0x6D, 0x1F)
PAGE_W_IN = 210 / 25.4
PAGE_H_IN = 297 / 25.4
MARGIN_IN = 1.0
USABLE_W_IN = PAGE_W_IN - MARGIN_IN * 2  # ~6.27in


def set_font(run, size=11, color=DARK, bold=False, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def spacer(doc, pt=18):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)
    p.paragraph_format.line_spacing = Pt(1)
    return p


def kicker(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, color=GREEN):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.upper())
    set_font(r, size=10.5, color=color, bold=True)
    r.font.underline = False
    # letter spacing via char spacing (approx)
    rPr = r._element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), "24")
    rPr.append(spacing)
    return p


def heading(doc, text, size=30, align=WD_ALIGN_PARAGRAPH.LEFT, color=DARK, space_after=10):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=True)
    return p


def body(doc, text, size=12.5, align=WD_ALIGN_PARAGRAPH.LEFT, color=MUTED, space_after=14, bold=False, line_spacing=1.4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    r = p.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    return p


def bullet(doc, text, color=DARK, size=11.5, marker="—"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.05)
    r1 = p.add_run(f"{marker}  ")
    set_font(r1, size=size, color=GREEN, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=size, color=color)
    return p


def check_item(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r1 = p.add_run("✓  ")
    set_font(r1, size=12.5, color=GREEN, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=12.5, color=DARK)
    return p


def cross_item(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r1 = p.add_run("–  ")
    set_font(r1, size=12.5, color=LIGHT_MUTED, bold=True)
    r2 = p.add_run(text)
    set_font(r2, size=12.5, color=MUTED)
    return p


def caption(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=10, color=LIGHT_MUTED, italic=True)
    return p


def image_full(doc, filename, max_width_in=USABLE_W_IN, center=True):
    path = os.path.join(ASSETS, filename)
    with Image.open(path) as im:
        w, h = im.size
    ratio = h / w
    width_in = max_width_in
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    return width_in * ratio


def image_inline(doc, filename, width_in, center=True):
    path = os.path.join(ASSETS, filename)
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    return p


def page_break(doc):
    doc.add_page_break()


def hr(doc, color=RGBColor(0xE5, 0xE7, 0xEB)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E5E7EB")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def module_page(doc, kicker_text, title, image_filename, description, features, badge_text=None):
    if badge_text:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(f"  {badge_text}  ")
        set_font(r, size=9.5, color=WARN_BG_TEXT, bold=True)
    kicker(doc, kicker_text)
    heading(doc, title, size=30, space_after=14)
    image_full(doc, image_filename)
    spacer(doc, 16)
    body(doc, description, size=13, color=DARK, space_after=14, line_spacing=1.35)
    for f in features:
        bullet(doc, f)


# ---------------------------------------------------------------------------
doc = Document()

section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Inches(MARGIN_IN)
section.bottom_margin = Inches(MARGIN_IN)
section.left_margin = Inches(MARGIN_IN)
section.right_margin = Inches(MARGIN_IN)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)

style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(11)
style.font.color.rgb = DARK
style.paragraph_format.line_spacing = 1.3
style.paragraph_format.space_after = Pt(8)

# footer with page number (muted, minimal)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.space_before = Pt(0)
run = fp.add_run()
set_font(run, size=9, color=LIGHT_MUTED)
fldChar1 = OxmlElement("w:fldChar")
fldChar1.set(qn("w:fldCharType"), "begin")
instrText = OxmlElement("w:instrText")
instrText.set(qn("xml:space"), "preserve")
instrText.text = "PAGE"
fldChar2 = OxmlElement("w:fldChar")
fldChar2.set(qn("w:fldCharType"), "end")
run._r.append(fldChar1)
run._r.append(instrText)
run._r.append(fldChar2)

# =============================================================================
# PORTADA
# =============================================================================
spacer(doc, 30)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(os.path.join(ASSETS, "logo.png"), width=Inches(0.9))

spacer(doc, 14)
heading(doc, "German CRM", size=44, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("CRM DE GESTIÓN COMERCIAL")
set_font(r, size=13, color=GREEN, bold=True)
rPr = r._element.get_or_add_rPr()
spacing = OxmlElement("w:spacing")
spacing.set(qn("w:val"), "30")
rPr.append(spacing)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(36)
r = p.add_run("La herramienta para centralizar clientes, productos,\npresupuestos y análisis comercial.")
set_font(r, size=14.5, color=MUTED)

spacer(doc, 10)
image_full(doc, "dashboard.png", max_width_in=USABLE_W_IN)

page_break(doc)

# =============================================================================
# INTRODUCCIÓN
# =============================================================================
spacer(doc, 60)
kicker(doc, "Introducción", align=WD_ALIGN_PARAGRAPH.CENTER)
heading(doc, "Una sola plataforma\npara toda la operación comercial", size=28, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
body(
    doc,
    "Centralizar toda la operación comercial de la empresa en una única plataforma, "
    "simplificando la gestión de clientes, productos, presupuestos y la toma de decisiones.",
    size=14.5,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=DARK,
    line_spacing=1.5,
)

page_break(doc)

# =============================================================================
# FUNCIONALIDADES — MÓDULOS
# =============================================================================
module_page(
    doc,
    "Funcionalidades",
    "Dashboard",
    "dashboard.png",
    "El Dashboard ofrece una visión rápida del estado general del negocio.",
    [
        "Clientes registrados",
        "Productos en catálogo",
        "Presupuestos generados",
        "Alertas de stock bajo",
        "Presupuestos y clientes recientes",
    ],
)
page_break(doc)

module_page(
    doc,
    "Funcionalidades",
    "Clientes",
    "clientes.png",
    "Permite administrar toda la cartera de clientes en un mismo lugar.",
    ["Alta y baja de clientes", "Modificación de datos", "Búsqueda instantánea", "Historial de contacto"],
)
page_break(doc)

module_page(
    doc,
    "Funcionalidades",
    "Productos",
    "productos.png",
    "Permite administrar el catálogo de productos con precios y stock actualizados.",
    ["Precio por producto", "Control de stock", "Categorías", "Alertas de stock bajo"],
)
page_break(doc)

module_page(
    doc,
    "Funcionalidades",
    "Presupuestos",
    "presupuesto-detalle.png",
    "Permite generar presupuestos profesionales en pocos segundos.",
    [
        "Selección de cliente",
        "Catálogo de productos",
        "Cantidades y precios editables",
        "Cálculo automático de totales",
        "Exportación a PDF",
    ],
)
page_break(doc)

module_page(
    doc,
    "Funcionalidades",
    "Competencia / Mercado",
    "competencia.png",
    "Permite registrar manualmente precios del mercado y compararlos con los propios.",
    ["Historial de precios", "Comparación directa", "Precio promedio", "Precio mínimo y máximo"],
)
page_break(doc)

module_page(
    doc,
    "Funcionalidades",
    "Proveedores",
    "proveedores-comparador2.png",
    "Permite gestionar proveedores y comparar costos de importación en tiempo real.",
    [
        "Gestión de proveedores",
        "Costos FOB",
        "Nacionalización",
        "Comparador de proveedores",
        "Historial de costos",
    ],
)
page_break(doc)

# =============================================================================
# DISEÑO
# =============================================================================
kicker(doc, "Diseño")
heading(doc, "Un sistema visual coherente", size=28, space_after=10)
body(doc, "Cards, botones y tablas con un mismo lenguaje visual. Diseño moderno, minimalista y 100% responsive.", size=13, color=DARK, space_after=20)

image_inline(doc, "crop-card.png", width_in=2.6)
spacer(doc, 6)
image_inline(doc, "crop-button.png", width_in=2.0)
spacer(doc, 6)
image_inline(doc, "crop-table.png", width_in=USABLE_W_IN)
spacer(doc, 10)
image_inline(doc, "mobile-clientes.png", width_in=2.0)
caption(doc, "Misma experiencia, en cualquier dispositivo.")

page_break(doc)

# =============================================================================
# ARQUITECTURA
# =============================================================================
spacer(doc, 60)
kicker(doc, "Arquitectura", align=WD_ALIGN_PARAGRAPH.CENTER)
heading(doc, "Construido para crecer", size=28, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

for title, desc in [
    ("Frontend moderno", "Interfaz rápida, fluida y preparada para cualquier dispositivo."),
    ("Base de datos escalable", "Estructura preparada para crecer junto con el negocio."),
    ("Arquitectura desacoplada", "Cada capa del sistema puede evolucionar de forma independiente."),
    ("Código escalable", "Base sólida, ordenada y lista para incorporar nuevas funcionalidades."),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, size=14, color=DARK, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(22)
    r2 = p2.add_run(desc)
    set_font(r2, size=12, color=MUTED)

page_break(doc)

# =============================================================================
# ALCANCE DEL PROYECTO
# =============================================================================
kicker(doc, "Alcance del proyecto")
heading(doc, "Qué incluye el desarrollo", size=28, space_after=22)

p = doc.add_paragraph()
r = p.add_run("INCLUYE")
set_font(r, size=11, color=GREEN, bold=True)
p.paragraph_format.space_after = Pt(10)

for item in [
    "Dashboard",
    "Clientes",
    "Productos",
    "Presupuestos",
    "Mercado",
    "Proveedores",
    "Base de datos",
    "Exportación a PDF",
    "Diseño responsive",
    "Diseño moderno",
    "Código fuente completo",
]:
    check_item(doc, item)

spacer(doc, 20)
p = doc.add_paragraph()
r = p.add_run("NO INCLUYE")
set_font(r, size=11, color=LIGHT_MUTED, bold=True)
p.paragraph_format.space_after = Pt(10)

for item in [
    "Integraciones externas",
    "Facturación electrónica",
    "Automatizaciones",
    "APIs",
    "Scraping",
    "Sincronización con servicios externos",
]:
    cross_item(doc, item)

spacer(doc, 6)
body(doc, "Estas funcionalidades podrán incorporarse en futuras etapas.", size=11, color=LIGHT_MUTED)

page_break(doc)

# =============================================================================
# ENTREGA
# =============================================================================
spacer(doc, 60)
kicker(doc, "Entrega", align=WD_ALIGN_PARAGRAPH.CENTER)
heading(doc, "Qué vas a recibir", size=28, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=26)

for item in ["Código fuente completo", "Base de datos", "Manual de instalación", "Documentación básica"]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(item)
    set_font(r, size=15, color=DARK)

page_break(doc)

# =============================================================================
# INVERSIÓN
# =============================================================================
spacer(doc, 70)
kicker(doc, "Inversión", align=WD_ALIGN_PARAGRAPH.CENTER)
heading(doc, "Desarrollo", size=24, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("USD 100")
set_font(r, size=52, color=GREEN, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(30)
r = p.add_run("Pago único")
set_font(r, size=13, color=MUTED)

body(
    doc,
    "El código fuente será propiedad del cliente.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=DARK,
    size=13,
)

spacer(doc, 40)
hr(doc)
spacer(doc, 30)

kicker(doc, "Plan de soporte continuo", align=WD_ALIGN_PARAGRAPH.CENTER)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("USD 30 / mes")
set_font(r, size=30, color=DARK, bold=True)

for item in [
    "Soporte técnico",
    "Corrección de errores",
    "Copias de seguridad",
    "Actualizaciones de compatibilidad",
    "Hasta 2 horas mensuales de pequeños ajustes o mejoras",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run("✓  ")
    set_font(r1, size=12.5, color=GREEN, bold=True)
    r2 = p.add_run(item)
    set_font(r2, size=12.5, color=DARK)

page_break(doc)

# =============================================================================
# EVOLUCIÓN DEL SISTEMA
# =============================================================================
spacer(doc, 40)
kicker(doc, "Evolución del sistema", align=WD_ALIGN_PARAGRAPH.CENTER)
heading(doc, "Pensado para seguir creciendo", size=26, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=22)

body(
    doc,
    "El sistema fue desarrollado con una arquitectura modular y escalable, permitiendo "
    "incorporar nuevas funcionalidades a medida que evolucionen las necesidades del negocio.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=DARK,
    size=13,
    line_spacing=1.5,
)

body(
    doc,
    "Las mejoras, nuevos módulos o modificaciones que excedan el alcance del desarrollo "
    "inicial no están incluidas dentro del plan de soporte mensual y serán presupuestadas "
    "de forma independiente.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=MUTED,
    size=12.5,
    line_spacing=1.5,
)

spacer(doc, 10)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
r = p.add_run("Valor de desarrollo adicional")
set_font(r, size=12, color=MUTED)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(22)
r = p.add_run("USD 30 / hora")
set_font(r, size=24, color=GREEN, bold=True)

body(
    doc,
    "Antes de comenzar cualquier desarrollo adicional, PRAGMA STUDIO presentará una "
    "estimación de horas y alcance para su aprobación. Ningún trabajo será realizado "
    "sin la conformidad previa del cliente.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=DARK,
    size=12.5,
    line_spacing=1.5,
)

spacer(doc, 20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Ejemplos de mejoras futuras")
set_font(r, size=11, color=LIGHT_MUTED, bold=True)

for item in [
    "Nuevos módulos",
    "Integraciones con otros sistemas",
    "Automatizaciones",
    "Reportes personalizados",
    "Adaptaciones a nuevos procesos internos",
    "Cambios en la lógica de negocio",
    "Nuevas funcionalidades solicitadas por el cliente",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(item)
    set_font(r, size=11.5, color=MUTED)

page_break(doc)

# =============================================================================
# CIERRE
# =============================================================================
spacer(doc, 110)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(os.path.join(ASSETS, "logo.png"), width=Inches(0.7))

spacer(doc, 26)
heading(doc, "Gracias por considerar\nnuestra propuesta.", size=26, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
body(
    doc,
    "Estamos convencidos de que este sistema permitirá optimizar la gestión comercial "
    "y acompañar el crecimiento de la empresa.",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    color=MUTED,
    size=13,
    line_spacing=1.5,
)

spacer(doc, 50)
hr(doc)
spacer(doc, 20)

for line, sz, col, bold in [
    ("PRAGMA STUDIO", 13, DARK, True),
    ("pragmasolucionesdigitales@gmail.com", 11.5, MUTED, False),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(line)
    set_font(r, size=sz, color=col, bold=bold)

doc.save(OUT)
print("Saved:", OUT)
